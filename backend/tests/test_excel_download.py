import os
import sys
import io
import pytest
import openpyxl

# Add backend to path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))
from app import app, issue_token

@pytest.fixture
def client():
    app.config['TESTING'] = True
    with app.test_client() as c:
        yield c

def test_download_requires_auth(client):
    res = client.get('/api/download?period=2026-08')
    assert res.status_code == 401

def test_download_different_periods(client):
    token = issue_token('Javohir', 'admin')
    headers = {'Authorization': f'Bearer {token}'}

    # August
    res_aug = client.get('/api/download?period=2026-08', headers=headers)
    assert res_aug.status_code == 200
    wb_aug = openpyxl.load_workbook(io.BytesIO(res_aug.data))
    ws_aug = wb_aug['Stansiyalar Hisoboti']
    aug_title = ws_aug['A1'].value
    aug_sum = ws_aug[f'E{ws_aug.max_row}'].value

    # January
    res_jan = client.get('/api/download?period=2026-01', headers=headers)
    assert res_jan.status_code == 200
    wb_jan = openpyxl.load_workbook(io.BytesIO(res_jan.data))
    ws_jan = wb_jan['Stansiyalar Hisoboti']
    jan_title = ws_jan['A1'].value
    jan_sum = ws_jan[f'E{ws_jan.max_row}'].value

    assert 'Avgust' in aug_title
    assert 'Yanvar' in jan_title
    assert aug_title != jan_title
    assert aug_sum != jan_sum

def test_download_ytd_and_all_sheets(client):
    token = issue_token('Javohir', 'admin')
    headers = {'Authorization': f'Bearer {token}'}

    res_ytd = client.get('/api/download?period=ytd', headers=headers)
    assert res_ytd.status_code == 200
    wb_ytd = openpyxl.load_workbook(io.BytesIO(res_ytd.data))
    assert 'Stansiyalar Hisoboti' in wb_ytd.sheetnames
    assert 'Oylar Tahlili' in wb_ytd.sheetnames
    assert 'Kunlik Trend' in wb_ytd.sheetnames

def test_download_regional_scoping(client):
    token = issue_token('samarqandkiosk', 'kiosk', region='samarqandkiosk@railway.uz')
    headers = {'Authorization': f'Bearer {token}'}

    res = client.get('/api/download?period=2026-08', headers=headers)
    assert res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res.data))
    ws = wb['Stansiyalar Hisoboti']
    assert 'Самарқанд' in ws['A1'].value
    # Max row should be 6 (title, subtitle, blank, header, 1 station, total)
    assert ws.max_row == 6
