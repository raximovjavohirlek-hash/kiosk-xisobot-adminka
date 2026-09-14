import os
import io
import sqlite3
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EMAIL_MAP_COLS = {
    'andijonkiosk@railway.uz': ('Андижон', 4, 5),
    'buxorokiosk@railway.uz': ('Бухоро', 6, 7),
    'khivakiosk@railway.uz': ('Хива', 8, 9),
    'kiosk@axonlogic.uz': ('Тошкент Жанубий', 10, 11),
    'margilonkiosk@railway.uz': ('Марғилон', 12, 13),
    'navoiykiosk@railway.uz': ('Навои', 14, 15),
    'namangankiosk@railway.uz': ('Наманган', 16, 17),
    'nukuskiosk@railway.uz': ('Нукус', 18, 19),
    'qarshikiosk@railway.uz': ('Қарши', 20, 21),
    'qongirotkiosk@railway.uz': ('Қўнғирод', 22, 23),
    'qoqonkiosk@railway.uz': ('Қўқон', 24, 25),
    'samarqandkiosk@railway.uz': ('Самарқанд', 26, 27),
    'termizkiosk@railway.uz': ('Термиз', 28, 29),
    'toshkent.shimoliykiosk@railway.uz': ('Тошкент Марказий', 30, 31),
    'urganchkiosk@railway.uz': ('Урганч', 32, 33)
}

OYLAR_EMAIL_COLS = {
    'toshkent.shimoliykiosk@railway.uz': (4, 5),
    'kiosk@axonlogic.uz': (6, 7),
    'samarqandkiosk@railway.uz': (8, 9),
    'navoiykiosk@railway.uz': (10, 11),
    'buxorokiosk@railway.uz': (12, 13),
    'nukuskiosk@railway.uz': (14, 15),
    'urganchkiosk@railway.uz': (16, 17),
    'khivakiosk@railway.uz': (18, 19),
    'qarshikiosk@railway.uz': (20, 21),
    'termizkiosk@railway.uz': (22, 23),
    'qongirotkiosk@railway.uz': (24, 25),
    'andijonkiosk@railway.uz': (26, 27),
    'qoqonkiosk@railway.uz': (28, 29),
    'margilonkiosk@railway.uz': (30, 31),
    'namangankiosk@railway.uz': (32, 33)
}

MONTH_NAMES_UZ = {
    1: ('Yanvar', 'Январь', 31),
    2: ('Fevral', 'Февраль', 28),
    3: ('Mart', 'Март', 31),
    4: ('Aprel', 'Апрель', 30),
    5: ('May', 'Май', 31),
    6: ('Iyun', 'Июнь', 30),
    7: ('Iyul', 'Июль', 31),
    8: ('Avgust', 'Август', 31),
    9: ('Sentabr', 'Сентябрь', 30),
    10: ('Oktabr', 'Октябрь', 31),
    11: ('Noyabr', 'Ноябрь', 30),
    12: ('Dekabr', 'Декабрь', 31)
}


def find_template_path():
    candidates = [
        os.path.join(os.path.dirname(__file__), 'excellar', 'Август кисока.xlsx'),
        os.path.join(os.path.dirname(__file__), 'Август кисока.xlsx'),
        os.path.join(os.path.dirname(os.path.dirname(__file__)), 'backend', 'excellar', 'Август кисока.xlsx'),
        os.path.join(os.path.dirname(os.path.dirname(__file__)), 'backend', 'Август кисока.xlsx'),
    ]
    for p in candidates:
        if os.path.exists(p):
            return p
    return None


def generate_kiosk_excel_report(db_path, period='latest', region=None, email_map=None):
    """
    Generates an Excel workbook with the EXACT layout, styling, and charts of 'Август кисока.xlsx':
    - Sheet 'Лист1': Ranked tables (summa & soni) and 2 BarCharts
    - Sheet 'Худудлар': 15 Stations x 31 Days matrix with totals formulas
    - Sheet 'Жами': Daily total, Online and Terminal breakdown with formulas
    - Sheet 'Ойлар кесимида': All months (Jan..Dec) x 15 stations
    - Sheet 'Stansiyalar Hisoboti': Executive summary table
    """
    template_path = find_template_path()
    if not template_path:
        raise FileNotFoundError("Reference template 'Август кисока.xlsx' not found.")

    wb = openpyxl.load_workbook(template_path)

    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    c = conn.cursor()

    # Determine target month & year
    target_year = 2026
    target_month = 9

    if period and period not in ('ytd', 'all', 'overall', 'latest'):
        parts = period.split('-')
        if len(parts) == 2:
            try:
                target_year = int(parts[0])
                target_month = int(parts[1])
            except ValueError:
                target_year, target_month = 2026, 9
    elif period == 'latest':
        # Find latest available month in DB
        c.execute("SELECT ym FROM tickets WHERE ym LIKE '2026-%' GROUP BY ym ORDER BY ym DESC LIMIT 1")
        row = c.fetchone()
        if row and row['ym']:
            parts = row['ym'].split('-')
            target_year, target_month = int(parts[0]), int(parts[1])

    target_ym = f"{target_year}-{target_month:02d}"
    month_info = MONTH_NAMES_UZ.get(target_month, ('Oy', 'Ой', 30))
    latin_month_name, uz_month_name, days_in_month = month_info

    # Fetch daily station numbers for the target month
    c.execute('''
        SELECT date_str, LOWER(TRIM(user_email)) as em, SUM(qty) as tickets, SUM(summa) as summa
        FROM tickets
        WHERE ym = ?
        GROUP BY date_str, LOWER(TRIM(user_email))
    ''', (target_ym,))
    daily_station_data = {}
    for r in c.fetchall():
        d_str = r['date_str']
        em = r['em']
        if d_str not in daily_station_data:
            daily_station_data[d_str] = {}
        daily_station_data[d_str][em] = (int(r['tickets'] or 0), round(float(r['summa'] or 0)))

    # Fetch daily payment breakdown for the target month
    c.execute('''
        SELECT date_str,
               SUM(CASE WHEN payment_type = 'Online' THEN qty ELSE 0 END) as on_t,
               SUM(CASE WHEN payment_type = 'Online' THEN summa ELSE 0 END) as on_s,
               SUM(CASE WHEN payment_type = 'Terminal' THEN qty ELSE 0 END) as term_t,
               SUM(CASE WHEN payment_type = 'Terminal' THEN summa ELSE 0 END) as term_s
        FROM tickets
        WHERE ym = ?
        GROUP BY date_str
    ''', (target_ym,))
    daily_pay_data = {}
    for r in c.fetchall():
        daily_pay_data[r['date_str']] = (
            int(r['on_t'] or 0), round(float(r['on_s'] or 0)),
            int(r['term_t'] or 0), round(float(r['term_s'] or 0))
        )

    # 1. Sheet: Худудлар
    if 'Худудлар' in wb.sheetnames:
        ws_hudud = wb['Худудлар']
        ws_hudud['D1'] = f"Чиптаҳоналардаги киоскалардан сотилган\n чипталар тўғрисида маълумотнома 01.{target_month:02d}.{target_year}-{days_in_month:02d}.{target_month:02d}.{target_year}. йил."

        for day in range(1, 32):
            r = 3 + day
            if day <= days_in_month:
                d_str = f"{day:02d}.{target_month:02d}.{target_year}"
                dt_val = datetime.datetime(target_year, target_month, day, 0, 0)
                ws_hudud.cell(r, 1).value = dt_val

                # Formulas for Total Tickets (Col B) and Total Summa (Col C) across all 15 stations
                ws_hudud.cell(r, 2).value = f"=D{r}+F{r}+H{r}+J{r}+L{r}+N{r}+P{r}+R{r}+T{r}+V{r}+X{r}+Z{r}+AB{r}+AD{r}+AF{r}"
                ws_hudud.cell(r, 3).value = f"=E{r}+G{r}+I{r}+K{r}+M{r}+O{r}+Q{r}+S{r}+U{r}+W{r}+Y{r}+AA{r}+AC{r}+AE{r}+AG{r}"

                day_stations = daily_station_data.get(d_str, {})
                has_day_data = d_str in daily_station_data

                for email, (st_name, col_soni, col_summa) in EMAIL_MAP_COLS.items():
                    em_lower = email.lower()
                    val_t, val_s = day_stations.get(em_lower, (0, 0))
                    if not val_t and not val_s and 'samarqand' in em_lower:
                        val_t, val_s = day_stations.get('samarkandkiosk@railway.uz', (0, 0))
                    if not val_t and not val_s and 'toshkent.shimoliy' in em_lower:
                        val_t, val_s = day_stations.get('tashkentkiosk@railway.uz', (0, 0))

                    if has_day_data:
                        ws_hudud.cell(r, col_soni).value = val_t
                        ws_hudud.cell(r, col_summa).value = val_s
                    else:
                        ws_hudud.cell(r, col_soni).value = None
                        ws_hudud.cell(r, col_summa).value = None
            else:
                # Clear extra day cells (e.g. day 31 for 30-day months)
                ws_hudud.cell(r, 1).value = None
                for c_idx in range(2, 34):
                    ws_hudud.cell(r, c_idx).value = None

        # Row 35 totals
        ws_hudud.cell(35, 2).value = "=SUM(B4:B34)"
        ws_hudud.cell(35, 3).value = "=SUM(C4:C34)"
        for c_idx in range(4, 34):
            col_letter = get_column_letter(c_idx)
            ws_hudud.cell(35, c_idx).value = f"=SUM({col_letter}4:{col_letter}34)"

    # 2. Sheet: Жами
    if 'Жами' in wb.sheetnames:
        ws_jami = wb['Жами']
        ws_jami['A1'] = f"Чиптаҳоналардаги киоскалардан сотилган\n чипталар тўғрисида маълумотнома 01.{target_month:02d}.{target_year}-{days_in_month:02d}.{target_month:02d}.{target_year}. йил."

        for day in range(1, 32):
            r = 3 + day
            if day <= days_in_month:
                d_str = f"{day:02d}.{target_month:02d}.{target_year}"
                dt_val = datetime.datetime(target_year, target_month, day, 0, 0)
                ws_jami.cell(r, 1).value = dt_val

                ws_jami.cell(r, 2).value = f"=Худудлар!B{r}"
                ws_jami.cell(r, 3).value = f"=Худудлар!C{r}"
                ws_jami.cell(r, 4).value = f"=B{r}-F{r}"
                ws_jami.cell(r, 5).value = f"=C{r}-G{r}"

                if d_str in daily_pay_data:
                    on_t, on_s, term_t, term_s = daily_pay_data[d_str]
                    ws_jami.cell(r, 6).value = term_t
                    ws_jami.cell(r, 7).value = term_s
                else:
                    ws_jami.cell(r, 6).value = None
                    ws_jami.cell(r, 7).value = None
            else:
                ws_jami.cell(r, 1).value = None
                for c_idx in range(2, 8):
                    ws_jami.cell(r, c_idx).value = None

        ws_jami.cell(35, 2).value = "=SUM(B4:B34)"
        ws_jami.cell(35, 3).value = "=SUM(C4:C34)"
        ws_jami.cell(35, 4).value = "=SUM(D4:D34)"
        ws_jami.cell(35, 5).value = "=SUM(E4:E34)"
        ws_jami.cell(35, 6).value = "=SUM(F4:F34)"
        ws_jami.cell(35, 7).value = "=SUM(G4:G34)"

    # 3. Sheet: Лист1 (Rankings & Charts)
    if 'Лист1' in wb.sheetnames:
        ws_list1 = wb['Лист1']
        ws_list1.cell(2, 1).value = f"Инфокиоскалардан 1-{uz_month_name} {days_in_month}-{uz_month_name}кунига қадар сотилган чипталар\n тўғрисида маълумотнома (худудлар кесимида (сумма))"
        ws_list1.cell(25, 1).value = f"Инфокиоскалардан  1-{uz_month_name} {days_in_month}-{uz_month_name} кунига қадар сотилган чипталар\n тўғрисида маълумотнома (худудлар кесимида (сотилган чипталар сони)) "

        station_sums = []
        for email, (st_name, col_soni, col_summa) in EMAIL_MAP_COLS.items():
            if period in ('ytd', 'all'):
                c.execute('''
                    SELECT SUM(qty), SUM(summa)
                    FROM tickets
                    WHERE ym LIKE '2026-%' AND (LOWER(TRIM(user_email)) = ? OR LOWER(TRIM(user_email)) = ?)
                ''', (email.lower(), email.lower().replace('samarqand', 'samarkand')))
            else:
                c.execute('''
                    SELECT SUM(qty), SUM(summa)
                    FROM tickets
                    WHERE ym = ? AND (LOWER(TRIM(user_email)) = ? OR LOWER(TRIM(user_email)) = ?)
                ''', (target_ym, email.lower(), email.lower().replace('samarqand', 'samarkand')))
            row = c.fetchone()
            soni = int(row[0] or 0)
            summa = round(float(row[1] or 0))
            station_sums.append({
                'station': st_name,
                'soni': soni,
                'summa': summa,
                'col_soni_letter': get_column_letter(col_soni),
                'col_summa_letter': get_column_letter(col_summa)
            })

        # Table 1: Sorted by Summa descending
        by_summa = sorted(station_sums, key=lambda x: x['summa'], reverse=True)
        for i, item in enumerate(by_summa, start=4):
            ws_list1.cell(i, 1).value = i - 3
            ws_list1.cell(i, 2).value = item['station']
            ws_list1.cell(i, 3).value = f"=Худудлар!{item['col_soni_letter']}35"
            ws_list1.cell(i, 4).value = f"=Худудлар!{item['col_summa_letter']}35"
        ws_list1.cell(20, 3).value = "=SUM(C4:C18)"
        ws_list1.cell(20, 4).value = "=SUM(D4:D18)"

        # Table 2: Sorted by Soni descending
        by_soni = sorted(station_sums, key=lambda x: x['soni'], reverse=True)
        for i, item in enumerate(by_soni, start=27):
            ws_list1.cell(i, 1).value = i - 26
            ws_list1.cell(i, 2).value = item['station']
            ws_list1.cell(i, 3).value = f"=Худудлар!{item['col_soni_letter']}35"
            ws_list1.cell(i, 4).value = f"=Худудлар!{item['col_summa_letter']}35"
        ws_list1.cell(43, 3).value = "=SUM(C27:C41)"
        ws_list1.cell(43, 4).value = "=SUM(D27:D41)"

    # 4. Sheet: Ойлар кесимида
    sheet_oylar_name = 'Ойлар кесимида' if 'Ойлар кесимида' in wb.sheetnames else ('Ойlar кесимида' if 'Ойlar кесимида' in wb.sheetnames else None)
    if sheet_oylar_name:
        ws_oylar = wb[sheet_oylar_name]
        c.execute('''
            SELECT ym, LOWER(TRIM(user_email)) as em, SUM(qty) as tickets, SUM(summa) as summa
            FROM tickets
            WHERE ym LIKE '2026-%'
            GROUP BY ym, LOWER(TRIM(user_email))
        ''')
        monthly_station_totals = {}
        for row in c.fetchall():
            ym_val = row['ym']
            if ym_val not in monthly_station_totals:
                monthly_station_totals[ym_val] = {}
            monthly_station_totals[ym_val][row['em']] = (int(row['tickets'] or 0), round(float(row['summa'] or 0)))

        for m_num in range(1, 13):
            r = 3 + m_num
            ym_str = f"2026-{m_num:02d}"
            m_dict = monthly_station_totals.get(ym_str)
            if m_dict:
                ws_oylar.cell(r, 2).value = f"=D{r}+F{r}+H{r}+J{r}+L{r}+N{r}+P{r}+R{r}+T{r}+V{r}+X{r}+Z{r}+AB{r}+AD{r}+AF{r}"
                ws_oylar.cell(r, 3).value = f"=E{r}+G{r}+I{r}+K{r}+M{r}+O{r}+Q{r}+S{r}+U{r}+W{r}+Y{r}+AA{r}+AC{r}+AE{r}+AG{r}"
                for email, (col_t, col_s) in OYLAR_EMAIL_COLS.items():
                    t_val, s_val = m_dict.get(email, (0, 0))
                    if not t_val and not s_val and 'samarqand' in email:
                        t_val, s_val = m_dict.get('samarkandkiosk@railway.uz', (0, 0))
                    if not t_val and not s_val and 'toshkent.shimoliy' in email:
                        t_val, s_val = m_dict.get('tashkentkiosk@railway.uz', (0, 0))
                    ws_oylar.cell(r, col_t).value = t_val
                    ws_oylar.cell(r, col_s).value = s_val

        ws_oylar.cell(16, 2).value = "=SUM(B4:B15)"
        ws_oylar.cell(16, 3).value = "=SUM(C4:C15)"
        for c_idx in range(4, 34):
            col_letter = get_column_letter(c_idx)
            ws_oylar.cell(16, c_idx).value = f"=SUM({col_letter}4:{col_letter}15)"

    # 5. Sheet: Stansiyalar Hisoboti (Executive summary table for backwards-compatibility & tests)
    ws_summary = wb.create_sheet(title="Stansiyalar Hisoboti")
    ws_summary.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Arial", size=10)
    bold_font = Font(name="Arial", size=10, bold=True)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    if period in ('ytd',):
        period_title_uz = f"{target_year} YTD (Yil boshidan beri)"
    elif period in ('all', 'overall'):
        period_title_uz = "Barcha Oylar Birgalikda"
    else:
        period_title_uz = f"{latin_month_name} ({uz_month_name}) {target_year} oyi"

    if region:
        st_meta = email_map.get(region.lower()) if email_map else None
        st_title = st_meta.get('station', region) if isinstance(st_meta, dict) else (EMAIL_MAP_COLS.get(region.lower(), (region, 0, 0))[0])
        summary_title = f"{st_title} — Kiosk Chipta Sotuvi Hisoboti ({period_title_uz})"
    else:
        summary_title = f"Kiosklar Bo'yicha Chipta Sotuvi Hisoboti ({period_title_uz})"

    ws_summary.merge_cells('A1:G1')
    ws_summary['A1'] = summary_title
    ws_summary['A1'].font = Font(name="Arial", size=14, bold=True, color="0F172A")
    ws_summary['A1'].alignment = Alignment(horizontal="center", vertical="center")
    ws_summary.row_dimensions[1].height = 30

    ws_summary.merge_cells('A2:G2')
    ws_summary['A2'] = f"O'zbekiston Temir Yo'llari — Kiosk Analytics PRO | Shakllantirilgan vaqt: {datetime.datetime.now().strftime('%d.%m.%Y %H:%M')}"
    ws_summary['A2'].font = Font(name="Arial", size=9, italic=True, color="64748B")
    ws_summary['A2'].alignment = Alignment(horizontal="center", vertical="center")
    ws_summary.row_dimensions[2].height = 18

    headers = ['№', 'Kassa Stansiyasi', 'Pochta Manzili', 'Chiptalar Soni (ta)', 'Tushum Summasi (so\'m)', 'Ulushi (%)', 'O\'rtacha Narx (so\'m)']
    ws_summary.append([])
    ws_summary.append(headers)
    ws_summary.row_dimensions[4].height = 25

    for col_num, h in enumerate(headers, 1):
        cell = ws_summary.cell(row=4, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Get station summary data
    if region:
        r_em = region.strip().lower()
        if period in ('ytd', 'all'):
            c.execute('''
                SELECT LOWER(TRIM(user_email)) as em, station_name, SUM(qty) as tickets, SUM(summa) as summa
                FROM tickets
                WHERE ym LIKE '2026-%' AND LOWER(TRIM(user_email)) = ?
                GROUP BY LOWER(TRIM(user_email))
            ''', (r_em,))
        else:
            c.execute('''
                SELECT LOWER(TRIM(user_email)) as em, station_name, SUM(qty) as tickets, SUM(summa) as summa
                FROM tickets
                WHERE ym = ? AND LOWER(TRIM(user_email)) = ?
                GROUP BY LOWER(TRIM(user_email))
            ''', (target_ym, r_em))
    else:
        if period in ('ytd', 'all'):
            c.execute('''
                SELECT LOWER(TRIM(user_email)) as em, station_name, SUM(qty) as tickets, SUM(summa) as summa
                FROM tickets
                WHERE ym LIKE '2026-%'
                GROUP BY LOWER(TRIM(user_email))
                ORDER BY summa DESC
            ''')
        else:
            c.execute('''
                SELECT LOWER(TRIM(user_email)) as em, station_name, SUM(qty) as tickets, SUM(summa) as summa
                FROM tickets
                WHERE ym = ?
                GROUP BY LOWER(TRIM(user_email))
                ORDER BY summa DESC
            ''', (target_ym,))

    sum_rows = c.fetchall()
    tot_tix = sum(int(r['tickets'] or 0) for r in sum_rows)
    tot_sum = sum(round(float(r['summa'] or 0)) for r in sum_rows) or 1

    r_idx = 5
    for idx, r in enumerate(sum_rows, 1):
        st_em = r['em']
        st_name = r['station_name'] or st_em
        t_val = int(r['tickets'] or 0)
        s_val = round(float(r['summa'] or 0))
        sh_pct = round((s_val / tot_sum * 100), 1)
        avg_p = round(s_val / t_val) if t_val > 0 else 0

        ws_summary.append([idx, st_name, st_em, t_val, s_val, sh_pct, avg_p])
        ws_summary.cell(row=r_idx, column=1).alignment = Alignment(horizontal="center")
        ws_summary.cell(row=r_idx, column=4).number_format = '#,##0'
        ws_summary.cell(row=r_idx, column=5).number_format = '#,##0'
        ws_summary.cell(row=r_idx, column=6).number_format = '0.0'
        ws_summary.cell(row=r_idx, column=7).number_format = '#,##0'
        for c_i in range(1, 8):
            cell = ws_summary.cell(row=r_idx, column=c_i)
            cell.font = data_font
            cell.border = thin_border
        r_idx += 1

    ws_summary.append(['', 'JAMI', '', tot_tix, tot_sum, 100.0, round(tot_sum / tot_tix) if tot_tix > 0 else 0])
    for c_i in range(1, 8):
        cell = ws_summary.cell(row=r_idx, column=c_i)
        cell.font = bold_font
        cell.border = thin_border
        cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    ws_summary.cell(row=r_idx, column=4).number_format = '#,##0'
    ws_summary.cell(row=r_idx, column=5).number_format = '#,##0'
    ws_summary.cell(row=r_idx, column=6).number_format = '0.0'
    ws_summary.cell(row=r_idx, column=7).number_format = '#,##0'

    for col in ws_summary.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_summary.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # If YTD/All, also add Oylar Tahlili & Kunlik Trend for complete analysis
    if period in ('ytd', 'all'):
        ws_m_tahlil = wb.create_sheet(title="Oylar Tahlili")
        ws_m_tahlil.views.sheetView[0].showGridLines = True
        m_headers = ['№', 'Hisobot Oyi', 'Chiptalar Soni (ta)', 'Tushum Summasi (so\'m)', 'Online Chiptalar', 'Online Summa', 'Terminal Chiptalar', 'Terminal Summa', 'Ulushi (%)']
        ws_m_tahlil.append(m_headers)
        for col_num, h in enumerate(m_headers, 1):
            cell = ws_m_tahlil.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        c.execute('''
            SELECT ym,
                   SUM(qty) as tix,
                   SUM(summa) as sum,
                   SUM(CASE WHEN payment_type = 'Online' THEN qty ELSE 0 END) as on_tix,
                   SUM(CASE WHEN payment_type = 'Online' THEN summa ELSE 0 END) as on_sum,
                   SUM(CASE WHEN payment_type = 'Terminal' THEN qty ELSE 0 END) as term_tix,
                   SUM(CASE WHEN payment_type = 'Terminal' THEN summa ELSE 0 END) as term_sum
            FROM tickets
            WHERE ym LIKE '2026-%'
            GROUP BY ym
            ORDER BY ym ASC
        ''')
        ym_rows = c.fetchall()
        for idx, ymr in enumerate(ym_rows, 1):
            m_code = ymr['ym']
            m_num = int(m_code.split('-')[1])
            m_name = MONTH_NAMES_UZ.get(m_num, (m_code, 30))[0] + ' 2026'
            tix = int(ymr['tix'] or 0)
            summ = round(float(ymr['sum'] or 0))
            on_t = int(ymr['on_tix'] or 0)
            on_s = round(float(ymr['on_sum'] or 0))
            term_t = int(ymr['term_tix'] or 0)
            term_s = round(float(ymr['term_sum'] or 0))
            sh_pct = round(summ / tot_sum * 100, 1)
            ws_m_tahlil.append([idx, m_name, tix, summ, on_t, on_s, term_t, term_s, sh_pct])

        ws_d_trend = wb.create_sheet(title="Kunlik Trend")
        ws_d_trend.views.sheetView[0].showGridLines = True
        d_headers = ['Sana', 'Jami Chiptalar (ta)', 'Jami Summa (so\'m)', 'Uzcard Terminal (ta)', 'Uzcard Summa', 'Humo Terminal (ta)', 'Humo Summa', 'Jami Terminal (ta)', 'Jami Terminal Summa', 'Online Chiptalar (ta)', 'Online Summa']
        ws_d_trend.append(d_headers)
        for col_num, h in enumerate(d_headers, 1):
            cell = ws_d_trend.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        c.execute('''
            SELECT date_str,
                   SUM(qty) as tix,
                   SUM(summa) as sum,
                   SUM(CASE WHEN payment_method LIKE '%Uzcard%' THEN qty ELSE 0 END) as uz_t,
                   SUM(CASE WHEN payment_method LIKE '%Uzcard%' THEN summa ELSE 0 END) as uz_s,
                   SUM(CASE WHEN (payment_method LIKE '%Humo%' OR payment_method LIKE '%Uzkassa%') THEN qty ELSE 0 END) as hu_t,
                   SUM(CASE WHEN (payment_method LIKE '%Humo%' OR payment_method LIKE '%Uzkassa%') THEN summa ELSE 0 END) as hu_s,
                   SUM(CASE WHEN payment_type = 'Terminal' THEN qty ELSE 0 END) as term_t,
                   SUM(CASE WHEN payment_type = 'Terminal' THEN summa ELSE 0 END) as term_s,
                   SUM(CASE WHEN payment_type = 'Online' THEN qty ELSE 0 END) as on_t,
                   SUM(CASE WHEN payment_type = 'Online' THEN summa ELSE 0 END) as on_s
            FROM tickets
            WHERE ym LIKE '2026-%'
            GROUP BY date_str
            ORDER BY date_str ASC
        ''')
        for dr in c.fetchall():
            ws_d_trend.append([
                dr['date_str'], int(dr['tix'] or 0), round(float(dr['sum'] or 0)),
                int(dr['uz_t'] or 0), round(float(dr['uz_s'] or 0)),
                int(dr['hu_t'] or 0), round(float(dr['hu_s'] or 0)),
                int(dr['term_t'] or 0), round(float(dr['term_s'] or 0)),
                int(dr['on_t'] or 0), round(float(dr['on_s'] or 0))
            ])

    # If scoped to single station (regional login)
    if region:
        # If regional user, keep only Stansiyalar Hisoboti and remove sheets of other stations
        for sname in ['Худудлар', 'Жами', 'Лист1', 'Ойлар кесимида', 'Oylar Tahlili', 'Kunlik Trend']:
            if sname in wb.sheetnames:
                del wb[sname]
        wb.active = wb['Stansiyalar Hisoboti']
    else:
        # Default active sheet is 'Лист1' with the 2 Charts and Rankings!
        if 'Лист1' in wb.sheetnames:
            wb.active = wb['Лист1']

    conn.close()

    out_buf = io.BytesIO()
    wb.save(out_buf)
    out_buf.seek(0)

    # Dynamic filename matching user expectations
    if region:
        st_meta = email_map.get(region.lower()) if email_map else None
        st_title = st_meta.get('station', region) if isinstance(st_meta, dict) else (EMAIL_MAP_COLS.get(region.lower(), (region, 0, 0))[0])
        safe_st = "".join(c for c in st_title if c.isalnum() or c in (' ', '_', '-')).strip()
        filename = f"Kiosk_Hisobot_{safe_st}_{target_ym}.xlsx"
    else:
        filename = f"{uz_month_name}_кисока.xlsx" if period not in ('ytd', 'all') else f"2026_{period.upper()}_кисока.xlsx"

    return out_buf, filename
